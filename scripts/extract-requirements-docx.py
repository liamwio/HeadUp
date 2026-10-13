"""Read-only structural extraction for the retained HeadUp requirements DOCX."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: extract-requirements-docx.py INPUT.docx")
    path = Path(sys.argv[1])
    document = Document(path)
    paragraphs = [
        {"index": index, "style": paragraph.style.name, "text": paragraph.text.strip()}
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip()
    ]
    tables = []
    for table_index, table in enumerate(document.tables):
        tables.append({
            "index": table_index,
            "rows": [
                ["\n".join(part.strip() for part in cell.text.splitlines() if part.strip()) for cell in row.cells]
                for row in table.rows
            ],
        })
    sections = []
    for section_index, section in enumerate(document.sections):
        sections.append({
            "index": section_index,
            "header": [paragraph.text.strip() for paragraph in section.header.paragraphs if paragraph.text.strip()],
            "footer": [paragraph.text.strip() for paragraph in section.footer.paragraphs if paragraph.text.strip()],
            "pageWidth": section.page_width,
            "pageHeight": section.page_height,
            "margins": {
                "top": section.top_margin,
                "right": section.right_margin,
                "bottom": section.bottom_margin,
                "left": section.left_margin,
            },
        })
    result = {
        "path": str(path),
        "paragraphCount": len(document.paragraphs),
        "nonemptyParagraphCount": len(paragraphs),
        "tableCount": len(tables),
        "sectionCount": len(sections),
        "paragraphs": paragraphs,
        "tables": tables,
        "sections": sections,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
